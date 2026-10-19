"""
Comprehensive Word Document (.docx) Generator for Final Project Report
Strictly adheres to Vellore Institute of Technology (VIT) Guidelines:
- A4 Paper Size (210 x 297 mm)
- Margins: Left: 1.5", Right: 1.0", Top: 1.0", Bottom: 1.0"
- Times New Roman throughout
- 1.5 line spacing for body, fully justified
- Headings: 16pt Uppercase Chapter, 14pt Uppercase H1, 12pt Title Case H2, 12pt Bold Italic H3
- Formatted tables with shaded headers
- High-res embedded figures with 10pt captions below
- Section-based centered page numbering: None for Prelim, Roman for Front Matter, Arabic for Chapters 1-8
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_page_number_field(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def configure_section_page(section, top=1.0, bottom=1.0, left=1.5, right=1.0):
    section.page_width = Inches(8.27)  # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)

def build_word_document():
    doc = Document()
    
    # Base Normal Style Configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # =========================================================================
    # SECTION 1: COVER & PRELIMINARY PAGES (No page numbering)
    # =========================================================================
    sec1 = doc.sections[0]
    configure_section_page(sec1)
    sec1.footer.is_linked_to_previous = False
    
    def add_title_page():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(24)
        r = p.add_run("<< CSI1006 – Course Code / Title >>")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(16)
        r.font.bold = True

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(36)
        r_title = p_title.add_run("OPTIMIZATION OF BATTERY MANAGEMENT SYSTEMS (BMS) IN ELECTRIC VEHICLES WITH NOVEL HGWO-PSO SENSOR SELECTION AND SAFETY-PRIORITIZED CRITICAL FAULT DETECTION")
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(15)
        r_title.font.bold = True

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(8)
        r_sub = p_sub.add_run("Submitted in partial fulfillment of the requirements for the degree of\n")
        r_sub.font.name = 'Times New Roman'
        r_sub.font.size = Pt(13)
        r_sub.font.italic = True
        
        r_deg = p_sub.add_run("MASTER OF TECHNOLOGY (INTEGRATED)\n")
        r_deg.font.name = 'Times New Roman'
        r_deg.font.size = Pt(16)
        r_deg.font.bold = True
        
        r_in = p_sub.add_run("in\n")
        r_in.font.italic = True
        r_in.font.size = Pt(13)
        
        r_prog = p_sub.add_run("<< Programme Name >>\n")
        r_prog.font.name = 'Times New Roman'
        r_prog.font.size = Pt(15)
        r_prog.font.bold = True

        p_by = doc.add_paragraph()
        p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_by.paragraph_format.space_before = Pt(18)
        p_by.paragraph_format.space_after = Pt(12)
        r_by = p_by.add_run("by")
        r_by.font.name = 'Times New Roman'
        r_by.font.size = Pt(13)
        r_by.font.italic = True

        # Students Table
        tbl = doc.add_table(rows=4, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        students = [
            ("<< Reg. No. 1 >>", "<< STUDENT NAME 1 >>"),
            ("<< Reg. No. 2 >>", "<< STUDENT NAME 2 >>"),
            ("<< Reg. No. 3 >>", "<< STUDENT NAME 3 >>"),
            ("<< Reg. No. 4 >>", "<< STUDENT NAME 4 >>"),
        ]
        for idx, (reg, name) in enumerate(students):
            row = tbl.rows[idx]
            row.cells[0].width = Inches(2.2)
            row.cells[1].width = Inches(3.2)
            p0 = row.cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r0 = p0.add_run(reg)
            r0.font.bold = True
            r0.font.size = Pt(12)
            p1 = row.cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r1 = p1.add_run(name)
            r1.font.bold = True
            r1.font.size = Pt(12)
            set_cell_margins(row.cells[0], top=40, bottom=40)
            set_cell_margins(row.cells[1], top=40, bottom=40)

        p_guide = doc.add_paragraph()
        p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_guide.paragraph_format.space_before = Pt(36)
        p_guide.paragraph_format.space_after = Pt(24)
        r_g1 = p_guide.add_run("Under the Supervision of\n")
        r_g1.font.size = Pt(13)
        r_g2 = p_guide.add_run("<< Project Guide Name >>\n")
        r_g2.font.bold = True
        r_g2.font.size = Pt(14)
        r_g3 = p_guide.add_run("<< Designation >>\n<< School Name >>\n")
        r_g3.font.size = Pt(12)

        p_date = doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_date.paragraph_format.space_before = Pt(30)
        r_date = p_date.add_run("<< Month Year >>")
        r_date.font.bold = True
        r_date.font.size = Pt(13)

    print("Adding Cover Page...")
    add_title_page()
    doc.add_page_break()

    print("Adding Front Page (Title Page replicate)...")
    add_title_page()
    doc.add_page_break()

    # Declaration
    print("Adding Declaration...")
    p_d_head = doc.add_paragraph()
    p_d_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_d_head.paragraph_format.space_before = Pt(36)
    p_d_head.paragraph_format.space_after = Pt(24)
    r_d_head = p_d_head.add_run("DECLARATION")
    r_d_head.font.size = Pt(16)
    r_d_head.font.bold = True

    p_d_body1 = doc.add_paragraph(
        'We hereby declare that the project entitled << "[Project Title]" >> submitted by us, '
        'for the award of the degree of << Degree >> in << Programme >> to VIT is a record of '
        'bonafide work carried out by us under the supervision of Prof. << Guide Name >>.'
    )
    p_d_body2 = doc.add_paragraph(
        'We further declare that the work reported in this project has not been submitted and '
        'will not be submitted, either in part or in full, for the award of any other degree or '
        'diploma in this institute or any other institute or university.'
    )
    p_d_place = doc.add_paragraph()
    p_d_place.paragraph_format.space_before = Pt(36)
    p_d_place.add_run("Place : Vellore\nDate  : << DD/MM/YYYY >>")

    p_d_sig = doc.add_paragraph()
    p_d_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_d_sig.paragraph_format.space_before = Pt(12)
    p_d_sig.add_run(
        "Signature of the Candidates\n\n"
        "<< Reg. No. 1 – STUDENT NAME 1 >>\n"
        "<< Reg. No. 2 – STUDENT NAME 2 >>\n"
        "<< Reg. No. 3 – STUDENT NAME 3 >>\n"
        "<< Reg. No. 4 – STUDENT NAME 4 >>"
    )
    doc.add_page_break()

    # Certificate
    print("Adding Certificate...")
    p_c_head = doc.add_paragraph()
    p_c_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c_head.paragraph_format.space_before = Pt(36)
    p_c_head.paragraph_format.space_after = Pt(24)
    r_c_head = p_c_head.add_run("CERTIFICATE")
    r_c_head.font.size = Pt(16)
    r_c_head.font.bold = True

    p_c_body1 = doc.add_paragraph(
        'This is to certify that the project entitled << "[Project Title]" >> submitted by '
        '<< Student Name (Reg. No.), Student Name (Reg. No.), Student Name (Reg. No.), Student Name (Reg. No.) >>, '
        '<< School Name >>, VIT, for the award of the degree of << Degree >> in << Programme >>, '
        'is a record of bonafide work carried out by the students under my supervision during '
        '<< Semester >>, as per the VIT code of academic and research ethics.'
    )
    p_c_body2 = doc.add_paragraph(
        'The contents of this report have not been submitted and will not be submitted either in '
        'part or in full, for the award of any other degree or diploma in this institute or any '
        'other institute or university. The project fulfills the requirements and regulations of '
        'the University and in my opinion meets the necessary standards for submission.'
    )
    p_c_place = doc.add_paragraph()
    p_c_place.paragraph_format.space_before = Pt(30)
    p_c_place.add_run("Place : Vellore\nDate  : << DD/MM/YYYY >>")

    p_c_sig = doc.add_paragraph()
    p_c_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_c_sig.paragraph_format.space_before = Pt(24)
    p_c_sig.add_run("Signature of the Guide\n\n\n")

    p_c_exam = doc.add_paragraph()
    p_c_exam.paragraph_format.space_before = Pt(24)
    p_c_exam.add_run("<< Internal Examiner – 1 >>" + " " * 45 + "<< Internal Examiner – 2 >>\n\n\n")
    
    p_c_hod = doc.add_paragraph()
    p_c_hod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c_hod.add_run("<< Head of the Department >>")
    doc.add_page_break()

    # Acknowledgement
    print("Adding Acknowledgement...")
    p_a_head = doc.add_paragraph()
    p_a_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a_head.paragraph_format.space_before = Pt(36)
    p_a_head.paragraph_format.space_after = Pt(24)
    r_a_head = p_a_head.add_run("ACKNOWLEDGEMENT")
    r_a_head.font.size = Pt(16)
    r_a_head.font.bold = True

    doc.add_paragraph(
        'We << Reg. No. – Name, Reg. No. – Name, Reg. No. – Name, Reg. No. – Name >> are deeply grateful '
        'to the management of Vellore Institute of Technology (VIT) for providing us with the opportunity '
        'and resources to undertake this project. Their commitment to fostering a conducive learning '
        'environment has been instrumental in our academic journey. The support and infrastructure provided '
        'by VIT have enabled us to explore and develop our ideas to their fullest potential.'
    )
    doc.add_paragraph(
        'Our sincere thanks to << Dean Name >>, the Dean of the << School Name >>, for constant support '
        'and encouragement. The Dean’s leadership and vision have greatly inspired us to strive for excellence. '
        'The dedication to academic excellence and innovation has been a constant source of motivation.'
    )
    doc.add_paragraph(
        'We express our profound appreciation to << Head of Department’s Name >>, the Head of the '
        '<< Department Name >>, for insightful guidance and continuous support. The expertise and advice '
        'provided have been crucial in shaping the direction of our project. The constructive feedback and '
        'encouragement have been invaluable in overcoming challenges and achieving our project goals.'
    )
    doc.add_paragraph(
        'We are immensely thankful to our project supervisor, << Supervisor’s Name >>, for dedicated '
        'mentorship and invaluable feedback. The patience, knowledge, and encouragement have been pivotal '
        'in the successful completion of this project.'
    )
    doc.add_paragraph(
        'We extend our heartfelt thanks to all for the guidance and support received throughout this endeavour.'
    )

    # =========================================================================
    # SECTION 2: FRONT MATTER (Executive Summary, TOC, Lists) - Roman Numerals (i, ii...)
    # =========================================================================
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section_page(sec2)
    sec2.footer.is_linked_to_previous = False
    
    # Configure Roman page numbering starting at i
    sectPr = sec2._sectPr
    pgNumType = parse_xml(r'<w:pgNumType %s w:fmt="romanLower" w:start="1"/>' % nsdecls('w'))
    sectPr.append(pgNumType)
    
    footer_p2 = sec2.footer.paragraphs[0]
    footer_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run2 = footer_p2.add_run()
    footer_run2.font.name = 'Times New Roman'
    footer_run2.font.size = Pt(11)
    add_page_number_field(footer_run2)

    # Read the updated markdown text
    with open("FINAL_PROJECT_REPORT_UPDATED.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    # Split between Front Matter and Chapters (Chapter 1 starts at "# CHAPTER 1")
    ch1_match = re.search(r'# CHAPTER 1[:\s]', md_text, re.IGNORECASE)
    if not ch1_match:
        ch1_match = re.search(r'# 1\.\s+INTRODUCTION', md_text, re.IGNORECASE)
    
    front_matter_md = md_text[:ch1_match.start()]
    body_md = md_text[ch1_match.start():]

    def render_markdown_section(text_chunk, is_body=False):
        lines = text_chunk.split("\n")
        in_table = False
        table_rows = []
        in_code = False
        code_lines = []
        
        i = 0
        while i < len(lines):
            line = lines[i]
            trimmed = line.strip()
            
            # Code blocks
            if trimmed.startswith("```"):
                if in_code:
                    # Flush code
                    code_text = "\n".join(code_lines)
                    p_code = doc.add_paragraph()
                    p_code.paragraph_format.line_spacing = 1.0
                    p_code.paragraph_format.space_before = Pt(4)
                    p_code.paragraph_format.space_after = Pt(6)
                    p_code.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    r_c = p_code.add_run(code_text)
                    r_c.font.name = 'Courier New'
                    r_c.font.size = Pt(9.5)
                    code_lines = []
                    in_code = False
                else:
                    in_code = True
                    code_lines = []
                i += 1
                continue
                
            if in_code:
                code_lines.append(line)
                i += 1
                continue

            # Tables
            if trimmed.startswith("|") and trimmed.endswith("|"):
                table_rows.append(trimmed)
                i += 1
                continue
            elif table_rows:
                # Flush table
                render_table_rows(table_rows)
                table_rows = []
                
            # Blank line
            if not trimmed:
                i += 1
                continue
                
            # Horizontal rule
            if trimmed in ["---", "***", "___"]:
                i += 1
                continue

            # Figure embed checks
            fig_match = re.search(r'\*Figure\s+(7\.[1-5]):\s*(.*?)\*', trimmed)
            if fig_match:
                fig_num = fig_match.group(1)
                fig_cap = fig_match.group(2)
                embed_figure(fig_num, fig_cap)
                i += 1
                continue

            # Headings
            if trimmed.startswith("# "):
                # Chapter level
                title = trimmed[2:].strip()
                if is_body and not title.startswith("CHAPTER 1"):
                    doc.add_page_break()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(12)
                p.paragraph_format.line_spacing = 1.5
                r = p.add_run(title.upper())
                r.font.name = 'Times New Roman'
                r.font.size = Pt(16)
                r.font.bold = True
                i += 1
                continue
                
            elif trimmed.startswith("## "):
                title = trimmed[3:].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(title.upper())
                r.font.name = 'Times New Roman'
                r.font.size = Pt(14)
                r.font.bold = True
                i += 1
                continue
                
            elif trimmed.startswith("### "):
                title = trimmed[4:].strip()
                # Table Title check
                if title.upper().startswith("TABLE "):
                    p_tbl = doc.add_paragraph()
                    p_tbl.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_tbl.paragraph_format.space_before = Pt(12)
                    p_tbl.paragraph_format.space_after = Pt(4)
                    r_t = p_tbl.add_run(title)
                    r_t.font.name = 'Times New Roman'
                    r_t.font.size = Pt(10.5)
                    r_t.font.bold = True
                else:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(4)
                    r = p.add_run(title)
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(12)
                    r.font.bold = True
                i += 1
                continue
                
            elif trimmed.startswith("#### "):
                title = trimmed[5:].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(title)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.italic = True
                i += 1
                continue

            # Regular Paragraph / List items
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)

            # Bullet or Numbered check
            if trimmed.startswith("- ") or trimmed.startswith("* "):
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run("• ")
                parse_markdown_formatting(trimmed[2:], p)
            elif re.match(r'^\d+\.\s+', trimmed):
                num_match = re.match(r'^(\d+\.)\s+', trimmed)
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run(num_match.group(1) + " ")
                parse_markdown_formatting(trimmed[len(num_match.group(0)):], p)
            else:
                parse_markdown_formatting(trimmed, p)

            i += 1

        if table_rows:
            render_table_rows(table_rows)

    def render_table_rows(table_lines):
        # Parse Markdown Table
        data = []
        for line in table_lines:
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                continue  # separator
            data.append(cells)
            
        if not data:
            return

        cols = len(data[0])
        tbl = doc.add_table(rows=len(data), cols=cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for r_idx, row_data in enumerate(data):
            row = tbl.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx in range(min(cols, len(row_data))):
                cell = row.cells[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                
                if is_header:
                    set_cell_background(cell, "1E3A8A")  # Deep blue header
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")  # Subtle alternating tint
                    
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (is_header or c_idx > 0) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                raw_cell_text = row_data[c_idx]
                # Format cell text
                clean_cell_text = raw_cell_text.replace("**", "").replace("`", "")
                r = p.add_run(clean_cell_text)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5 if is_header else 9.0)
                if is_header:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                elif "**" in raw_cell_text:
                    r.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def parse_markdown_formatting(raw_text, paragraph):
        # Clean math brackets
        clean = raw_text.replace("$", "").replace("`", "")
        # Tokenize bold text
        tokens = re.split(r'(\*\*.*?\*\*)', clean)
        for tok in tokens:
            if tok.startswith("**") and tok.endswith("**"):
                r = paragraph.add_run(tok[2:-2])
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.bold = True
            else:
                r = paragraph.add_run(tok)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)

    def embed_figure(fig_num, caption):
        img_map = {
            "7.1": "convergence_curves.png",
            "7.2": "cross_dataset_comparison.png",
            "7.3": "shap_summary_bar.png",
            "7.4": "shap_beeswarm_critical.png",
            "7.5": "shap_dependence_temp_soc.png"
        }
        img_file = img_map.get(fig_num)
        if img_file and os.path.exists(img_file):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            width = Inches(5.8) if fig_num in ["7.2", "7.1"] else Inches(5.2)
            p_img.add_run().add_picture(img_file, width=width)
            
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Figure {fig_num}: {caption}")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        r_cap.font.bold = True

    print("Rendering Front Matter (Executive Summary, TOC, Lists)...")
    render_markdown_section(front_matter_md, is_body=False)

    # =========================================================================
    # SECTION 3: CHAPTERS 1 TO 8 + REFERENCES + APPENDICES - Arabic Numerals (1, 2, 3...)
    # =========================================================================
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section_page(sec3)
    sec3.footer.is_linked_to_previous = False
    
    # Configure Arabic page numbering starting at 1
    sectPr3 = sec3._sectPr
    pgNumType3 = parse_xml(r'<w:pgNumType %s w:fmt="decimal" w:start="1"/>' % nsdecls('w'))
    sectPr3.append(pgNumType3)
    
    footer_p3 = sec3.footer.paragraphs[0]
    footer_p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run3 = footer_p3.add_run()
    footer_run3.font.name = 'Times New Roman'
    footer_run3.font.size = Pt(11)
    add_page_number_field(footer_run3)

    print("Rendering Body Chapters 1 through 8, References, and Appendices...")
    render_markdown_section(body_md, is_body=True)

    out_file = "FINAL_PROJECT_REPORT.docx"
    doc.save(out_file)
    print(f"Word Document successfully created: {out_file}")

    # Copy to user's Downloads directory for immediate convenience
    dl_out = os.path.expanduser("~/Downloads/FINAL_PROJECT_REPORT.docx")
    try:
        import shutil
        shutil.copyfile(out_file, dl_out)
        print(f"Also copied Word Document to: {dl_out}")
    except Exception as e:
        print(f"Notice copying to downloads: {e}")

if __name__ == '__main__':
    build_word_document()
