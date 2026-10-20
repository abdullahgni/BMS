"""
Build the finalized, perfectly formatted report for both:
1. ~/Downloads/final_project_report_v2.docx (and ~/Downloads/FINAL_PROJECT_REPORT.docx)
2. ~/Downloads/final_project_report_v2.md (and ~/Downloads/FINAL_PROJECT_REPORT.md)

Strictly adheres to Vellore Institute of Technology (VIT) guidelines:
- A4 Paper Size
- Margins: Left 1.5", Right 1.0", Top 1.0", Bottom 1.0"
- Times New Roman font throughout
- 1.5 line spacing for body text, fully justified
- Headings: 16pt Bold Uppercase Level 0, 14pt Bold Uppercase Level 1, 12pt Bold Title Case Level 2, 12pt Bold Italic Level 3
- Full Front Matter with Cover Page, Title Page, Declaration, Certificate, Acknowledgement, Exec Summary, Formatted TOC, List of Figures, List of Tables, Abbreviations, Symbols
- Formatted tables with Table caption placed ABOVE, deep-blue shaded header (#1E3A8A), clean borders
- High-resolution embedded figures with Figure caption placed BELOW, centered, 10pt font
- No broken ASCII art, no stray empty bullets, no raw LaTeX markup leakage
- Decoupled section numbering: Prelim (none), Front Matter (Roman i, ii...), Body (Arabic 1, 2, 3...)
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls

def clean_latex(text):
    if not text:
        return text
    
    # Direct string replacements (safe, fast, clean)
    repls = [
        ("$^\\circ\\text{C}$", "°C"),
        ("$^\\circ C$", "°C"),
        ("$^\\circ$", "°"),
        ("$\\text{m}\\Omega$", "mΩ"),
        ("$\\text{mg}$", "mg"),
        ("$\\text{kPa}$", "kPa"),
        ("$\\text{kW}$", "kW"),
        ("$\\text{Wh}$", "Wh"),
        ("$\\text{s}$", "s"),
        ("$\\text{A}$", "A"),
        ("$\\text{V}$", "V"),
        ("$\\text{ms}$", "ms"),
        ("$\\text{KB}$", "KB"),
        ("$\\text{MB}$", "MB"),
        ("$\\text{kg}$", "kg"),
        ("$\\ge 60\\%$", "≥ 60%"),
        ("$\\ge 12$", "≥ 12"),
        ("$\\approx 2.4\\text{ kg}$", "≈ 2.4 kg"),
        ("$\\approx 0.7\\text{ kg}$", "≈ 0.7 kg"),
        ("$< 10\\text{ ms}$", "< 10 ms"),
        ("$< 500\\text{ KB}$", "< 500 KB"),
        ("$14.2\\text{ ms}$", "14.2 ms"),
        ("$2.1\\text{ ms}$", "2.1 ms"),
        ("$1.85\\text{ MB}$", "1.85 MB"),
        ("$320\\text{ KB}$", "320 KB"),
        ("$\\alpha$", "α"),
        ("$\\beta$", "β"),
        ("$\\delta$", "δ"),
        ("$\\omega$", "ω"),
        ("$c_1, c_2$", "c₁, c₂"),
        ("$r_1, r_2$", "r₁, r₂"),
        ("$\\mathbf{x}_{feat}$", "x_feat"),
        ("$\\in \\{0, 1\\}^D$", "∈ {0, 1}^D"),
        ("$d_{sel}$", "d_sel"),
        ("$n_{est}$", "n_est"),
        ("$max\\_d$", "max_d"),
        ("$w_{crit}$", "w_crit"),
        ("$w_{sens}$", "w_sens"),
        ("$w_{0}=1.5$", "w₀ = 1.5"),
        ("$w_{0}$", "w₀"),
        ("$w_0$", "w₀"),
        ("$v_{i,d}^t$", "v_i,dᵗ"),
        ("$S(v)$", "S(v)"),
        ("$D$", "D"),
        ("$F$", "F"),
        ("$k=5$", "k = 5"),
        ("$k = 5$", "k = 5"),
        ("$N$", "N"),
        ("$t$", "t"),
        ("$T_{max}$", "T_max"),
        ("$|SHAP|$", "|SHAP|"),
        ("|\\text{SHAP}|", "|SHAP|"),
        ("$\\phi_j$", "φ_j"),
        ("$\\sigma$", "σ"),
        ("$\\mu$", "μ"),
        ("\\left(", "("),
        ("\\right)", ")"),
    ]
    res = text
    for old_s, new_s in repls:
        res = res.replace(old_s, new_s)
    res = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', res)
    res = re.sub(r'\\[a-zA-Z]+\{([^}]+)\}', r'\1', res)
    res = re.sub(r'\$([^\$]+)\$', r'\1', res)
    return res

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_page_number_field(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def configure_section_page(section, top=1.0, bottom=1.0, left=1.5, right=1.0):
    section.page_width = Inches(8.27)  # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)

def build_full_report():
    print("Reading user markdown content...")
    with open("FINAL_PROJECT_REPORT_USER.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    # Pre-clean TOC in markdown text to replace bullet list with formal table
    toc_formal_table = """# TABLE OF CONTENTS

| Chapter No. | Contents | Page No. |
| :---: | :--- | :---: |
| | **Executive Summary** | **i** |
| | **List of Figures** | **v** |
| | **List of Tables** | **vi** |
| | **List of Abbreviations** | **vii** |
| | **Symbols and Notations** | **ix** |
| **1.** | **INTRODUCTION** | **1** |
| | 1.1 Background | 1 |
| | 1.1.1 Machine Learning Based Fault Detection | 1 |
| | 1.1.2 Metaheuristic Feature Selection and Optimization | 2 |
| | 1.1.3 Functional Safety and Explainable Artificial Intelligence | 2 |
| | 1.2 Motivation | 3 |
| | 1.3 Scope of the Project | 3 |
| **2.** | **PROJECT DESCRIPTION AND GOALS** | **5** |
| | 2.1 Literature Review | 5 |
| | 2.1.1 Machine Learning and Metaheuristic Sensor Optimization in BMS | 5 |
| | 2.1.2 Deep Learning Architectures in Battery State Estimation and Prognostics | 6 |
| | 2.1.3 Class Imbalance, Cost-Sensitive Safety, and Explainable AI | 6 |
| | 2.1.4 Cross-Dataset Validation and Battery Aging Prognostics | 7 |
| | 2.2 Research Gap | 15 |
| | 2.3 Objectives | 16 |
| | 2.4 Problem Statement | 16 |
| | 2.5 Project Plan | 17 |
| **3.** | **TECHNICAL SPECIFICATION** | **19** |
| | 3.1 Requirements | 19 |
| | 3.1.1 Functional Requirements | 19 |
| | 3.1.2 Non-Functional Requirements | 19 |
| | 3.2 Feasibility Study | 21 |
| | 3.2.1 Technical Feasibility | 21 |
| | 3.2.2 Economic Feasibility | 21 |
| | 3.2.3 Social Feasibility | 21 |
| | 3.3 System Specification | 22 |
| | 3.3.1 Hardware Specification | 22 |
| | 3.3.2 Software Specification | 22 |
| **4.** | **SYSTEM DESIGN** | **23** |
| | 4.1 System Architecture | 23 |
| | 4.2 Detailed Design | 24 |
| | 4.2.1 Use Case Diagram and Description | 24 |
| | 4.2.2 Data Flow Diagram (DFD Level 0 & Level 1) | 25 |
| | 4.2.3 State Transition Diagram | 25 |
| | 4.2.4 Sequence Diagram | 25 |
| | 4.2.5 Class Diagram | 26 |
| | 4.2.6 Activity Diagram | 26 |
| **5.** | **METHODOLOGY AND TESTING** | **27** |
| | 5.1 Algorithm Formulation | 27 |
| | Algorithm 1: Proposed HGWO-PSO Feature Selection and Hyperparameter Co-Optimization | 27 |
| | Algorithm 2: Multi-Objective Safety Fitness Evaluation and Stratified K-Fold CV | 28 |
| | 5.2 Module Description | 29 |
| | 5.2.1 Telemetry Preprocessing and 3-Class Target Formulations | 29 |
| | 5.2.2 HGWO-PSO Evolutionary Optimization Engine | 29 |
| | 5.2.3 Cost-Sensitive Classifier Training and Cross-Validation | 29 |
| | 5.2.4 SHAP Electrochemical Interpretability Engine | 30 |
| | 5.3 Testing and Quality Assurance | 30 |
| | 5.3.1 Unit Testing | 30 |
| | 5.3.2 Integration Testing | 30 |
| | 5.3.3 Cross-Validation Testing | 31 |
| | 5.3.4 Cross-Dataset Generalization Testing | 31 |
| **6.** | **PROJECT IMPLEMENTATION** | **32** |
| | 6.1 Experimental Setup | 32 |
| | 6.2 Datasets with Description | 32 |
| | 6.2.1 EV Battery Fast-Charging Thermal Runaway Dataset | 32 |
| | 6.2.2 NASA Ames Li-ion Battery Prognostics Dataset | 34 |
| | 6.3 Step-by-Step Implementation Details | 35 |
| **7.** | **RESULTS AND DISCUSSION** | **37** |
| | 7.1 Baseline Reproduction and Benchmark Evaluation | 37 |
| | 7.2 Safety Critical Performance and Missed Detection Rate Analysis | 39 |
| | 7.3 Cross-Dataset Universal Generalization Findings | 39 |
| | 7.3.1 Benchmark Evaluation on Classic 4-Cell NASA Dataset | 39 |
| | 7.3.2 Multi-Temperature Global Archive Generalization (2,750 Cycles) | 40 |
| | 7.4 Explainable AI (SHAP) Electrochemical Interpretability | 41 |
| | 7.5 Physical Hardware and Engineering Cost Analysis | 43 |
| **8.** | **CONCLUSION AND FUTURE ENHANCEMENTS** | **45** |
| | 8.1 Conclusion | 45 |
| | 8.2 Future Enhancements | 45 |
| | **REFERENCES** | **47** |
| | **APPENDIX A – SAMPLE CODE** | **53** |
| | **APPENDIX B – CONFERENCE / JOURNAL PUBLICATION DETAILS** | **55** |"""

    # Replace bullet TOC with formal TOC table
    md_text = re.sub(r'# TABLE OF CONTENTS\n\n- \*\*Executive Summary\*\*.*?(?=# LIST OF FIGURES)', 
                     toc_formal_table + "\n\n---\n\n", md_text, flags=re.DOTALL)

    # Update List of Figures to include Page No column if missing
    lof_replacement = """# LIST OF FIGURES

| Figure No. | Title | Page No. |
| :---: | :--- | :---: |
| 4.1 | High-Level System Architecture of Proposed HGWO-PSO BMS Optimization Framework | 23 |
| 4.2 | Use Case Diagram for Intelligent BMS Telemetry Diagnostic System | 24 |
| 4.3 | Data Flow Diagram (Level 0 Context Diagram and Level 1 Processing Pipeline) | 25 |
| 4.4 | State Transition Diagram for Battery Fault Operating States | 25 |
| 4.5 | Sequence Diagram of HGWO-PSO Co-Optimization and Cross-Validation Execution | 25 |
| 4.6 | UML Class Diagram of the Proposed Machine Learning Optimization Architecture | 26 |
| 4.7 | Activity Diagram for Safety-Prioritized Fitness Evaluation and Sensor Reduction | 26 |
| 7.1 | Convergence Curves of BPSO, BWOA, BHPWOA, and Proposed HGWO-PSO | 37 |
| 7.2 | Cross-Dataset Performance Comparison on EV Thermal Runaway and NASA Datasets | 40 |
| 7.3 | Global Mean Absolute SHAP Importance Across Retained Battery Sensor Channels | 42 |
| 7.4 | SHAP Beeswarm Distribution for Critical Thermal Runaway Classification | 42 |
| 7.5 | Non-Linear SHAP Interaction Dependence Plot: Maximum Temperature vs. State of Charge | 43 |"""

    md_text = re.sub(r'# LIST OF FIGURES\n\n\| Figure No\. \| Title \|.*?(?=# LIST OF TABLES)',
                     lof_replacement + "\n\n---\n\n", md_text, flags=re.DOTALL)

    # Update List of Tables to include Page No column if missing
    lot_replacement = """# LIST OF TABLES

| Table No. | Title | Page No. |
| :---: | :--- | :---: |
| 2.1 | Comprehensive Synthesis of Reviewed Literature Across Machine Learning, Metaheuristics, and BMS Safety | 7 |
| 2.2 | Project Milestones, Deliverables, and Schedule Breakdown | 17 |
| 3.1 | Functional and Non-Functional Requirement Specifications | 20 |
| 3.2 | System Hardware and Development Software Specifications | 22 |
| 6.1 | Summary of Physical Telemetry Sensors in EV Fast-Charging Thermal Runaway Dataset | 32 |
| 6.2 | Attribute Description of NASA Ames Li-ion Battery Prognostics Dataset | 34 |
| 7.1 | Benchmark Comparison Across Baselines and Proposed HGWO-PSO on EV BMS Dataset | 37 |
| 7.2 | Active Physical Sensor Sets Retained Across Optimization Frameworks | 38 |
| 7.3 | Cross-Dataset Validation Results on NASA Li-ion Battery Prognostics Benchmark (Classic 4 Cells) | 39 |
| 7.4 | Multi-Temperature Global NASA Archive Benchmark (Strict 5-Fold Stratified CV, 2,750 Cycles) | 40 |
| 7.5 | Mean Absolute SHAP Attribute Ranking and Electrochemical Function in Battery Safety | 41 |
| 7.6 | Physical Hardware, Wiring Harness, and Computational Resource Reduction Metrics | 43 |"""

    md_text = re.sub(r'# LIST OF TABLES\n\n\| Table No\. \| Title \|.*?(?=# LIST OF ABBREVIATIONS)',
                     lot_replacement + "\n\n---\n\n", md_text, flags=re.DOTALL)

    # Save cleaned markdown
    out_md = "FINAL_PROJECT_REPORT_CLEAN.md"
    with open(out_md, "w", encoding="utf-8") as f:
        f.write(md_text)
    print(f"Saved clean markdown: {out_md}")

    # Build Word Document
    doc = Document()
    
    # Base Normal Style Configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # SECTION 1: COVER & PRELIMINARY PAGES (No page numbering)
    sec1 = doc.sections[0]
    configure_section_page(sec1)
    sec1.footer.is_linked_to_previous = False

    def add_title_page():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(24)
        r = p.add_run("<< CSI1006 – Course Code / Title >>")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(16)
        r.font.bold = True

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(36)
        r_title = p_title.add_run("OPTIMIZATION OF BATTERY MANAGEMENT SYSTEMS (BMS) IN ELECTRIC VEHICLES WITH NOVEL HGWO-PSO SENSOR SELECTION AND SAFETY-PRIORITIZED CRITICAL FAULT DETECTION")
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(15)
        r_title.font.bold = True

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(8)
        r_sub = p_sub.add_run("Submitted in partial fulfillment of the requirements for the degree of\n")
        r_sub.font.name = 'Times New Roman'
        r_sub.font.size = Pt(13)
        r_sub.font.italic = True
        
        r_deg = p_sub.add_run("MASTER OF TECHNOLOGY (INTEGRATED)\n")
        r_deg.font.name = 'Times New Roman'
        r_deg.font.size = Pt(16)
        r_deg.font.bold = True
        
        r_in = p_sub.add_run("in\n")
        r_in.font.italic = True
        r_in.font.size = Pt(13)
        
        r_prog = p_sub.add_run("<< Programme Name >>\n")
        r_prog.font.name = 'Times New Roman'
        r_prog.font.size = Pt(15)
        r_prog.font.bold = True

        p_by = doc.add_paragraph()
        p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_by.paragraph_format.space_before = Pt(18)
        p_by.paragraph_format.space_after = Pt(12)
        r_by = p_by.add_run("by")
        r_by.font.name = 'Times New Roman'
        r_by.font.size = Pt(13)
        r_by.font.italic = True

        # Students Table
        tbl = doc.add_table(rows=4, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        students = [
            ("<< Reg. No. 1 >>", "<< STUDENT NAME 1 >>"),
            ("<< Reg. No. 2 >>", "<< STUDENT NAME 2 >>"),
            ("<< Reg. No. 3 >>", "<< STUDENT NAME 3 >>"),
            ("<< Reg. No. 4 >>", "<< STUDENT NAME 4 >>"),
        ]
        for idx, (reg, name) in enumerate(students):
            row = tbl.rows[idx]
            row.cells[0].width = Inches(2.2)
            row.cells[1].width = Inches(3.2)
            p0 = row.cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r0 = p0.add_run(reg)
            r0.font.bold = True
            r0.font.size = Pt(12)
            p1 = row.cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r1 = p1.add_run(name)
            r1.font.bold = True
            r1.font.size = Pt(12)
            set_cell_margins(row.cells[0], top=40, bottom=40)
            set_cell_margins(row.cells[1], top=40, bottom=40)

        p_guide = doc.add_paragraph()
        p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_guide.paragraph_format.space_before = Pt(36)
        p_guide.paragraph_format.space_after = Pt(24)
        r_g1 = p_guide.add_run("Under the Supervision of\n")
        r_g1.font.size = Pt(13)
        r_g2 = p_guide.add_run("<< Project Guide Name >>\n")
        r_g2.font.bold = True
        r_g2.font.size = Pt(14)
        r_g3 = p_guide.add_run("<< Designation >>\n<< School Name >>\n")
        r_g3.font.size = Pt(12)

        p_date = doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_date.paragraph_format.space_before = Pt(30)
        r_date = p_date.add_run("<< Month Year >>")
        r_date.font.bold = True
        r_date.font.size = Pt(13)

    print("Adding Cover Page...")
    add_title_page()
    doc.add_page_break()

    print("Adding Front Page (Title Page replica)...")
    add_title_page()
    doc.add_page_break()

    # Declaration
    print("Adding Declaration...")
    p_d_head = doc.add_paragraph()
    p_d_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_d_head.paragraph_format.space_before = Pt(36)
    p_d_head.paragraph_format.space_after = Pt(24)
    r_d_head = p_d_head.add_run("DECLARATION")
    r_d_head.font.size = Pt(16)
    r_d_head.font.bold = True

    doc.add_paragraph(
        'We hereby declare that the project entitled << "[Project Title]" >> submitted by us, '
        'for the award of the degree of << Degree >> in << Programme >> to VIT is a record of '
        'bonafide work carried out by us under the supervision of Prof. << Guide Name >>.'
    )
    doc.add_paragraph(
        'We further declare that the work reported in this project has not been submitted and '
        'will not be submitted, either in part or in full, for the award of any other degree or '
        'diploma in this institute or any other institute or university.'
    )
    p_d_place = doc.add_paragraph()
    p_d_place.paragraph_format.space_before = Pt(36)
    p_d_place.add_run("Place : Vellore\nDate  : << DD/MM/YYYY >>")

    p_d_sig = doc.add_paragraph()
    p_d_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_d_sig.paragraph_format.space_before = Pt(12)
    p_d_sig.add_run(
        "Signature of the Candidates\n\n"
        "<< Reg. No. 1 – STUDENT NAME 1 >>\n"
        "<< Reg. No. 2 – STUDENT NAME 2 >>\n"
        "<< Reg. No. 3 – STUDENT NAME 3 >>\n"
        "<< Reg. No. 4 – STUDENT NAME 4 >>"
    )
    doc.add_page_break()

    # Certificate
    print("Adding Certificate...")
    p_c_head = doc.add_paragraph()
    p_c_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c_head.paragraph_format.space_before = Pt(36)
    p_c_head.paragraph_format.space_after = Pt(24)
    r_c_head = p_c_head.add_run("CERTIFICATE")
    r_c_head.font.size = Pt(16)
    r_c_head.font.bold = True

    doc.add_paragraph(
        'This is to certify that the project entitled << "[Project Title]" >> submitted by '
        '<< Student Name (Reg. No.), Student Name (Reg. No.), Student Name (Reg. No.), Student Name (Reg. No.) >>, '
        '<< School Name >>, VIT, for the award of the degree of << Degree >> in << Programme >>, '
        'is a record of bonafide work carried out by the students under my supervision during '
        '<< Semester >>, as per the VIT code of academic and research ethics.'
    )
    doc.add_paragraph(
        'The contents of this report have not been submitted and will not be submitted either in '
        'part or in full, for the award of any other degree or diploma in this institute or any '
        'other institute or university. The project fulfills the requirements and regulations of '
        'the University and in my opinion meets the necessary standards for submission.'
    )
    p_c_place = doc.add_paragraph()
    p_c_place.paragraph_format.space_before = Pt(30)
    p_c_place.add_run("Place : Vellore\nDate  : << DD/MM/YYYY >>")

    p_c_sig = doc.add_paragraph()
    p_c_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_c_sig.paragraph_format.space_before = Pt(24)
    p_c_sig.add_run("Signature of the Guide\n\n\n")

    p_c_exam = doc.add_paragraph()
    p_c_exam.paragraph_format.space_before = Pt(24)
    p_c_exam.add_run("<< Internal Examiner – 1 >>" + " " * 45 + "<< Internal Examiner – 2 >>\n\n\n")
    
    p_c_hod = doc.add_paragraph()
    p_c_hod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c_hod.add_run("<< Head of the Department >>")
    doc.add_page_break()

    # Acknowledgement
    print("Adding Acknowledgement...")
    p_a_head = doc.add_paragraph()
    p_a_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a_head.paragraph_format.space_before = Pt(36)
    p_a_head.paragraph_format.space_after = Pt(24)
    r_a_head = p_a_head.add_run("ACKNOWLEDGEMENT")
    r_a_head.font.size = Pt(16)
    r_a_head.font.bold = True

    doc.add_paragraph(
        'We << Reg. No. – Name, Reg. No. – Name, Reg. No. – Name, Reg. No. – Name >> are deeply grateful '
        'to the management of Vellore Institute of Technology (VIT) for providing us with the opportunity '
        'and resources to undertake this project. Their commitment to fostering a conducive learning '
        'environment has been instrumental in our academic journey. The support and infrastructure provided '
        'by VIT have enabled us to explore and develop our ideas to their fullest potential.'
    )
    doc.add_paragraph(
        'Our sincere thanks to << Dean Name >>, the Dean of the << School Name >>, for constant support '
        'and encouragement. The Dean’s leadership and vision have greatly inspired us to strive for excellence. '
        'The dedication to academic excellence and innovation has been a constant source of motivation.'
    )
    doc.add_paragraph(
        'We express our profound appreciation to << Head of Department’s Name >>, the Head of the '
        '<< Department Name >>, for insightful guidance and continuous support. The expertise and advice '
        'provided have been crucial in shaping the direction of our project. The constructive feedback and '
        'encouragement have been invaluable in overcoming challenges and achieving our project goals.'
    )
    doc.add_paragraph(
        'We are immensely thankful to our project supervisor, << Supervisor’s Name >>, for dedicated '
        'mentorship and invaluable feedback. The patience, knowledge, and encouragement have been pivotal '
        'in the successful completion of this project.'
    )
    doc.add_paragraph(
        'We extend our heartfelt thanks to all for the guidance and support received throughout this endeavour.'
    )

    # SECTION 2: FRONT MATTER (Roman numerals)
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section_page(sec2)
    sec2.footer.is_linked_to_previous = False
    
    sectPr2 = sec2._sectPr
    pgNumType2 = parse_xml(r'<w:pgNumType %s w:fmt="romanLower" w:start="1"/>' % nsdecls('w'))
    sectPr2.append(pgNumType2)
    
    footer_p2 = sec2.footer.paragraphs[0]
    footer_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run2 = footer_p2.add_run()
    footer_run2.font.name = 'Times New Roman'
    footer_run2.font.size = Pt(11)
    add_page_number_field(footer_run2)

    # Split markdown text into front matter and body chapters
    ch1_match = re.search(r'# CHAPTER 1[:\s]', md_text, re.IGNORECASE)
    front_matter_md = md_text[:ch1_match.start()]
    body_md = md_text[ch1_match.start():]

    def render_markdown(text_chunk, is_body=False):
        lines = text_chunk.split("\n")
        table_rows = []
        in_code = False
        code_lines = []
        
        i = 0
        while i < len(lines):
            line = lines[i]
            trimmed = line.strip()

            # Skip duplicate raw figure titles in text
            if re.match(r'^(Figure\s+\d+:\s*.*?)$', trimmed):
                i += 1
                continue

            # Code blocks
            if trimmed.startswith("```"):
                if in_code:
                    code_text = "\n".join(code_lines)
                    # Skip raw ASCII diagram text because high-res PNG images are embedded!
                    if "+--" in code_text or "│" in code_text or "▼" in code_text:
                        pass
                    else:
                        p_code = doc.add_paragraph()
                        p_code.paragraph_format.line_spacing = 1.0
                        p_code.paragraph_format.space_before = Pt(4)
                        p_code.paragraph_format.space_after = Pt(6)
                        p_code.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        r_c = p_code.add_run(code_text)
                        r_c.font.name = 'Courier New'
                        r_c.font.size = Pt(9.5)
                    code_lines = []
                    in_code = False
                else:
                    in_code = True
                    code_lines = []
                i += 1
                continue
                
            if in_code:
                code_lines.append(line)
                i += 1
                continue

            # Skip ASCII box lines
            if trimmed.startswith("+--") or (trimmed.startswith("|") and "+--" in trimmed):
                i += 1
                continue

            # Markdown Tables
            if trimmed.startswith("|") and trimmed.endswith("|"):
                table_rows.append(trimmed)
                i += 1
                continue
            elif table_rows:
                render_table_rows(table_rows)
                table_rows = []

            # Blank lines
            if not trimmed:
                i += 1
                continue

            # Horizontal rules
            if trimmed in ["---", "***", "___"]:
                i += 1
                continue

            # Figure Captions & Image Embedding
            fig_match = re.search(r'\*Figure\s+([47]\.[1-7]):\s*(.*?)\*', trimmed)
            if fig_match:
                fig_num = fig_match.group(1)
                fig_cap = fig_match.group(2)
                embed_figure(fig_num, fig_cap)
                i += 1
                continue

            # Level 0 Chapter Titles
            if trimmed.startswith("# "):
                title = trimmed[2:].strip()
                if is_body and not title.startswith("CHAPTER 1"):
                    doc.add_page_break()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(14)
                p.paragraph_format.line_spacing = 1.5
                r = p.add_run(clean_latex(title).upper())
                r.font.name = 'Times New Roman'
                r.font.size = Pt(16)
                r.font.bold = True
                i += 1
                continue

            # Level 1 Main Headings
            elif trimmed.startswith("## "):
                title = trimmed[3:].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(clean_latex(title).upper())
                r.font.name = 'Times New Roman'
                r.font.size = Pt(14)
                r.font.bold = True
                i += 1
                continue

            # Level 2 Sub-Headings
            elif trimmed.startswith("### "):
                title = trimmed[4:].strip()
                if title.upper().startswith("TABLE "):
                    p_tbl = doc.add_paragraph()
                    p_tbl.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_tbl.paragraph_format.space_before = Pt(14)
                    p_tbl.paragraph_format.space_after = Pt(4)
                    r_t = p_tbl.add_run(clean_latex(title))
                    r_t.font.name = 'Times New Roman'
                    r_t.font.size = Pt(10.5)
                    r_t.font.bold = True
                else:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(4)
                    r = p.add_run(clean_latex(title))
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(12)
                    r.font.bold = True
                i += 1
                continue

            # Level 3 Sub-sub-headings
            elif trimmed.startswith("#### "):
                title = trimmed[5:].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(clean_latex(title))
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.italic = True
                i += 1
                continue

            # Skip stray empty bullets or duplicate numbers
            if trimmed in ["•", "", "-", "*", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9."]:
                if i + 1 < len(lines) and lines[i+1].strip():
                    next_text = lines[i+1].strip()
                    p = doc.add_paragraph()
                    p.paragraph_format.line_spacing = 1.5
                    p.paragraph_format.space_after = Pt(4)
                    if trimmed in ["•", "", "-", "*"]:
                        p.paragraph_format.left_indent = Inches(0.25)
                        p.add_run("• ")
                        parse_styled_text(clean_latex(next_text), p)
                    else:
                        p.paragraph_format.left_indent = Inches(0.25)
                        p.add_run(trimmed + " ")
                        parse_styled_text(clean_latex(next_text), p)
                    i += 2
                    continue
                else:
                    i += 1
                    continue

            # Standard paragraph or list
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)

            # Bullet points
            if trimmed.startswith("- ") or trimmed.startswith("* "):
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run("• ")
                parse_styled_text(clean_latex(trimmed[2:]), p)
            elif re.match(r'^\d+\.\s+', trimmed):
                num_m = re.match(r'^(\d+\.)\s+', trimmed)
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run(num_m.group(1) + " ")
                parse_styled_text(clean_latex(trimmed[len(num_m.group(0)):]), p)
            else:
                parse_styled_text(clean_latex(trimmed), p)

            i += 1

        if table_rows:
            render_table_rows(table_rows)

    def render_table_rows(table_lines):
        data = []
        for line in table_lines:
            cells = [clean_latex(c.strip()) for c in line.strip("|").split("|")]
            if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                continue
            data.append(cells)
            
        if not data:
            return

        cols = len(data[0])
        tbl = doc.add_table(rows=len(data), cols=cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for r_idx, row_data in enumerate(data):
            row = tbl.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx in range(min(cols, len(row_data))):
                cell = row.cells[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                
                if is_header:
                    set_cell_background(cell, "1E3A8A")  # Deep Navy Header
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")  # Subtle alternating tint
                    
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (is_header or c_idx == 0 or c_idx == cols-1) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                cell_val = row_data[c_idx].replace("**", "").replace("`", "")
                r = p.add_run(cell_val)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5 if is_header else 9.0)
                if is_header:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                elif "**" in row_data[c_idx]:
                    r.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def parse_styled_text(text, paragraph):
        clean = text.replace("`", "")
        tokens = re.split(r'(\*\*.*?\*\*)', clean)
        for tok in tokens:
            if tok.startswith("**") and tok.endswith("**"):
                r = paragraph.add_run(tok[2:-2])
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.bold = True
            else:
                r = paragraph.add_run(tok)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)

    def embed_figure(fig_num, caption):
        img_map = {
            "4.1": "fig_4_1_architecture.png",
            "4.2": "fig_4_2_usecase.png",
            "4.3": "fig_4_3_dfd.png",
            "4.4": "fig_4_4_state_machine.png",
            "4.5": "fig_4_5_sequence.png",
            "4.6": "fig_4_6_class.png",
            "4.7": "fig_4_7_activity.png",
            "7.1": "convergence_curves.png",
            "7.2": "cross_dataset_comparison.png",
            "7.3": "shap_summary_bar.png",
            "7.4": "shap_beeswarm_critical.png",
            "7.5": "shap_dependence_temp_soc.png"
        }
        img_file = img_map.get(fig_num)
        if img_file and os.path.exists(img_file):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            width = Inches(5.8) if fig_num in ["4.1", "4.3", "4.5", "4.6", "7.1", "7.2"] else Inches(5.2)
            p_img.add_run().add_picture(img_file, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Figure {fig_num}: {caption}")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        r_cap.font.bold = True
        r_cap.font.italic = True

    print("Rendering Front Matter...")
    render_markdown(front_matter_md, is_body=False)

    # SECTION 3: CHAPTERS 1 TO 8 + REFERENCES + APPENDICES (Arabic Numerals)
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section_page(sec3)
    sec3.footer.is_linked_to_previous = False
    
    sectPr3 = sec3._sectPr
    pgNumType3 = parse_xml(r'<w:pgNumType %s w:fmt="decimal" w:start="1"/>' % nsdecls('w'))
    sectPr3.append(pgNumType3)
    
    footer_p3 = sec3.footer.paragraphs[0]
    footer_p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run3 = footer_p3.add_run()
    footer_run3.font.name = 'Times New Roman'
    footer_run3.font.size = Pt(11)
    add_page_number_field(footer_run3)

    print("Rendering Body Chapters 1 through 8, References, and Appendices...")
    render_markdown(body_md, is_body=True)

    out_docx = "final_project_report_v2.docx"
    doc.save(out_docx)
    print(f"Generated finalized Word Document: {out_docx}")

    # Copy files to Downloads
    import shutil
    shutil.copyfile(out_docx, os.path.expanduser("~/Downloads/final_project_report_v2.docx"))
    shutil.copyfile(out_docx, os.path.expanduser("~/Downloads/FINAL_PROJECT_REPORT.docx"))
    shutil.copyfile(out_md, os.path.expanduser("~/Downloads/final_project_report_v2.md"))
    shutil.copyfile(out_md, os.path.expanduser("~/Downloads/FINAL_PROJECT_REPORT.md"))
    print("All files updated in ~/Downloads/")

if __name__ == '__main__':
    build_full_report()
